import asyncio
import logging
import os
import re
import json
import traceback
from dotenv import load_dotenv
from aiogram import Bot, Dispatcher, types, F
from aiogram.filters import CommandStart, Command, StateFilter
from aiogram.types import (Message, FSInputFile, BotCommand,
                           InlineKeyboardMarkup, InlineKeyboardButton,
                           CallbackQuery, ReplyKeyboardMarkup, KeyboardButton)
from aiogram.fsm.state import State, StatesGroup
from aiogram.fsm.context import FSMContext
from aiogram.dispatcher.middlewares.base import BaseMiddleware
from typing import Callable, Dict, Any, Awaitable
from datetime import datetime, timedelta
import pandas as pd
import docx
import requests
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

import database

load_dotenv()
BOT_TOKEN = os.getenv("TELEGRAM_BOT_TOKEN")
SUPER_ADMIN_ID = int(os.getenv('SUPER_ADMIN_ID', 7631409423))
TRIAL_ACTIVATION_KEY = os.getenv('TRIAL_ACTIVATION_KEY', 'ABSYCODE-VIP-2026')

logging.basicConfig(level=logging.INFO)
bot = Bot(token=BOT_TOKEN)
dp = Dispatcher()

# ── AI OCR Configuration ─────────────────────────────────────────────────────
# EasyOCR and PyTorch dependencies have been removed to make the app lightweight.
# We now use Gemini 2.5 Flash native multimodal capabilities for OCR.

# database.init_db() is called in async main() below
os.makedirs('temp', exist_ok=True)

# ─────────────────────────────────────────────────────────────
# FSM STATES
# ─────────────────────────────────────────────────────────────
class MarineStates(StatesGroup):
    # Existing
    waiting_for_expense_details   = State()
    waiting_for_trip_note         = State()
    waiting_for_activation_key    = State()
    waiting_for_boat_name         = State()

    # Add Invoice
    waiting_for_invoice_category  = State()
    waiting_for_invoice_details   = State()
    # Edit Invoice
    waiting_for_edit_value        = State()
    # Invoice Image
    waiting_for_image_invoice_sel = State()
    waiting_for_image_upload      = State()
    # Delete single
    waiting_for_del_confirm       = State()
    # Admin panel
    admin_waiting_for_custom_expiry = State()
    # v4.0 — Voice correction (user edits transcript before injection)
    waiting_for_voice_correction  = State()
    # v4.0 — Report inline editing (price/text edit from detailed report)
    waiting_for_report_edit_value = State()
    # v4.0 — Merge target selection
    waiting_for_merge_target      = State()

pending_approval_requests: set = set()
pending_invoices: dict = {}

# ─────────────────────────────────────────────────────────────
# SUBSCRIPTION MIDDLEWARE
# ─────────────────────────────────────────────────────────────
class SubscriptionMiddleware(BaseMiddleware):
    async def __call__(self, handler: Callable, event: Any, data: Dict[str, Any]) -> Any:
        user = getattr(event, "from_user", None)
        if not user:
            return await handler(event, data)
        user_id = user.id
        if user_id == SUPER_ADMIN_ID:
            return await handler(event, data)
        state: FSMContext = data.get('state')
        if state:
            current_state = await state.get_state()
            if current_state in (MarineStates.waiting_for_activation_key.state, MarineStates.waiting_for_boat_name.state):
                return await handler(event, data)
        user_data = await database.get_user(user_id)
        is_active = False
        if user_data:
            is_active_db, expiry_str, _ = user_data
            if is_active_db and expiry_str:
                try:
                    if datetime.fromisoformat(expiry_str) > datetime.now():
                        is_active = True
                except ValueError:
                    pass
        if not is_active:
            if isinstance(event, CallbackQuery):
                await event.answer("⚠️ برجاء تفعيل الحساب أولاً لاستخدام هذه الميزة.", show_alert=True)
                return
            if isinstance(event, Message):
                if state:
                    await state.set_state(MarineStates.waiting_for_activation_key)
                if user_id not in pending_approval_requests:
                    pending_approval_requests.add(user_id)
                    kb = InlineKeyboardMarkup(inline_keyboard=[[
                        InlineKeyboardButton(text="✅ موافقة", callback_data=f"approve_{user_id}"),
                        InlineKeyboardButton(text="❌ رفض",    callback_data=f"reject_{user_id}")
                    ]])
                    try:
                        await event.bot.send_message(
                            SUPER_ADMIN_ID,
                            f"🚨 طلب انضمام جديد!\n👤 {user.full_name}\n🆔 {user_id}",
                            reply_markup=kb)
                    except Exception as e:
                        logging.error(f"Admin notify error: {e}")
                    await event.answer("تم إرسال طلب انضمامك للإدارة ⏳\nأو أدخل كود التفعيل:")
                else:
                    await event.answer("طلبك قيد المراجعة ⏳\nأو أدخل كود التفعيل:")
            return
        return await handler(event, data)

dp.message.middleware(SubscriptionMiddleware())
dp.callback_query.middleware(SubscriptionMiddleware())

# ─────────────────────────────────────────────────────────────
# MAIN MENU  (Req #6 — persistent, resize, all users)
# ─────────────────────────────────────────────────────────────
def get_main_menu() -> ReplyKeyboardMarkup:
    return ReplyKeyboardMarkup(
        keyboard=[
            [KeyboardButton(text="🛳️ مكتبي البحري"),   KeyboardButton(text="🌊 سجل السرحات")],
            [KeyboardButton(text="🛒 تجهيز رحلة"),      KeyboardButton(text="🚨 طوارئ وإنقاذ")],
            [KeyboardButton(text="💰 الحسابات"),         KeyboardButton(text="🌤️ حالة البحر")],
            [KeyboardButton(text="➕ إضافة فاتورة"),     KeyboardButton(text="✏️ تعديل فاتورة")],
            [KeyboardButton(text="🗑️ مسح"),              KeyboardButton(text="📸 إضافة صورة الفاتورة")],
            [KeyboardButton(text="📊 التقارير")]
        ],
        resize_keyboard=True,
        is_persistent=True
    )

# ─────────────────────────────────────────────────────────────
# CATEGORY HELPERS
# ─────────────────────────────────────────────────────────────
INVOICE_CATEGORIES = ["صيانة", "ماركت", "أدوات نظافة", "إكرامية", "عام", "بنزين"]

CATEGORY_EMOJIS = {
    "صيانة": "🔧", "بنزين": "⛽", "وقود": "⛽", "رواتب": "💵",
    "ماركت": "🛒", "إكرامية": "💝", "أدوات نظافة": "🧹",
    "طعام": "🍔", "أكل": "🍔", "مشتريات": "🛍️",
    "أدوات": "🛠️", "عام": "🏷️"
}

def get_category_emoji(category: str) -> str:
    if not category:
        return "🏷️"
    for key, emoji in CATEGORY_EMOJIS.items():
        if key in category:
            return emoji
    return "🏷️"

def get_category_keyboard() -> InlineKeyboardMarkup:
    buttons = []
    row = []
    for i, cat in enumerate(INVOICE_CATEGORIES):
        emoji = get_category_emoji(cat)
        row.append(InlineKeyboardButton(text=f"{emoji} {cat}", callback_data=f"inv_cat_{cat}"))
        if len(row) == 2:
            buttons.append(row)
            row = []
    if row:
        buttons.append(row)
    return InlineKeyboardMarkup(inline_keyboard=buttons)

# ─────────────────────────────────────────────────────────────
# /start
# ─────────────────────────────────────────────────────────────
@dp.message(CommandStart())
async def cmd_start(message: Message, state: FSMContext):
    await state.clear()
    await message.answer(
        "مرحباً بك في بوت AbsyCode Marine Assistant! ⚓\n\n"
        "أنا هنا لمساعدتك في إدارة مركبك ورحلاتك وحساباتك بكل سهولة.\n"
        "اختر من القائمة بالأسفل للبدء:",
        reply_markup=get_main_menu()
    )

# ─────────────────────────────────────────────────────────────
# REQ #2 — إضافة فاتورة  (Add Invoice)
# ─────────────────────────────────────────────────────────────
@dp.message(F.text == "➕ إضافة فاتورة")
async def handle_add_invoice_menu(message: Message, state: FSMContext):
    await state.clear()
    await state.set_state(MarineStates.waiting_for_invoice_category)
    await message.answer(
        "📂 اختر قسم الفاتورة:",
        reply_markup=get_category_keyboard()
    )

@dp.callback_query(MarineStates.waiting_for_invoice_category, F.data.startswith("inv_cat_"))
async def handle_invoice_category_selected(callback: CallbackQuery, state: FSMContext):
    category = callback.data.replace("inv_cat_", "")
    await state.update_data(invoice_category=category)
    data = await state.get_data()

    # Voice shortcut: if transcript exists, auto-save instead of waiting for text
    voice_text = data.get('voice_transcript')
    if voice_text:
        amount = _extract_amount(voice_text)
        user_data = await database.get_user(callback.from_user.id)
        boat_name = user_data[2] if user_data and user_data[2] else "مركب غير مسمى"
        invoice_id = await database.save_invoice(
            callback.from_user.id, amount, "جنيه", voice_text, category=category
        )
        await state.clear()
        emoji = get_category_emoji(category)
        continue_kb = InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="➕ إضافة فاتورة أخرى", callback_data="inv_add_another")],
            [InlineKeyboardButton(text="✅ إنهاء", callback_data="inv_finish")],
        ])
        await callback.message.edit_text(
            f"✅ **تم حفظ الفاتورة من الرسالة الصوتية!**\n\n"
            f"{emoji} القسم: {category}\n"
            f"💰 المبلغ: {amount:,.2f} جنيه\n"
            f"📝 التفاصيل: {voice_text}\n"
            f"🆔 رقم الفاتورة: #{invoice_id}\n\n"
            f"رحلة سعيدة يا ريس! 🛥️",
            reply_markup=continue_kb,
            parse_mode='Markdown'
        )
        await callback.answer()
        return

    # Normal manual flow: ask user to type details
    await state.set_state(MarineStates.waiting_for_invoice_details)
    emoji = get_category_emoji(category)
    await callback.message.edit_text(
        f"{emoji} القسم: **{category}**\n\n"
        "اكتب المبلغ وتفاصيل الفاتورة في رسالة واحدة.\n"
        "مثال: `500 تصليح مكينة`",
        parse_mode='Markdown'
    )
    await callback.answer()

@dp.message(MarineStates.waiting_for_invoice_details, F.text & ~F.text.startswith('/'))
async def handle_invoice_details_input(message: Message, state: FSMContext):
    data = await state.get_data()
    category = data.get('invoice_category', 'عام')
    text = message.text.strip()
    amount = _extract_amount(text)
    user_data = await database.get_user(message.from_user.id)
    boat_name = user_data[2] if user_data and user_data[2] else "مركب غير مسمى"
    invoice_id = await database.save_invoice(
        message.from_user.id, amount, "جنيه", text, category=category
    )
    emoji = get_category_emoji(category)
    continue_kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="➕ إضافة فاتورة أخرى", callback_data="inv_add_another")],
        [InlineKeyboardButton(text="✅ إنهاء", callback_data="inv_finish")],
    ])
    await message.answer(
        f"✅ **تم حفظ الفاتورة بنجاح!**\n\n"
        f"{emoji} القسم: {category}\n"
        f"💰 المبلغ: {amount:,.2f} جنيه\n"
        f"📝 التفاصيل: {text}\n"
        f"🆔 رقم الفاتورة: #{invoice_id}\n\n"
        f"رحلة سعيدة يا ريس! 🛥️",
        reply_markup=continue_kb,
        parse_mode='Markdown'
    )

@dp.callback_query(F.data == "inv_add_another")
async def handle_invoice_add_another(callback: CallbackQuery, state: FSMContext):
    await state.set_state(MarineStates.waiting_for_invoice_category)
    await callback.message.edit_text(
        "📂 اختر قسم الفاتورة:",
        reply_markup=get_category_keyboard()
    )
    await callback.answer()

@dp.callback_query(F.data == "inv_finish")
async def handle_invoice_finish(callback: CallbackQuery, state: FSMContext):
    await state.clear()
    await callback.message.edit_text("✅ تم الانتهاء من إدخال الفواتير. شغل جميل يا ريس! ⚓")
    await callback.message.answer("اختر من القائمة:", reply_markup=get_main_menu())
    await callback.answer()

# ─────────────────────────────────────────────────────────────
# REQ #3 — تعديل فاتورة  (Edit Invoice)
# ─────────────────────────────────────────────────────────────
EDIT_PAGE_SIZE = 5

@dp.message(F.text == "✏️ تعديل فاتورة")
async def handle_edit_invoice_menu(message: Message, state: FSMContext):
    await state.clear()
    await _show_invoice_list_for_edit(message, message.from_user.id, page=0)

async def _show_invoice_list_for_edit(target, user_id: int, page: int = 0):
    total = await database.count_user_invoices(user_id)
    is_bot_msg = False
    if hasattr(target, 'from_user') and target.from_user:
        is_bot_msg = target.from_user.is_bot

    if total == 0:
        text = "لا توجد فواتير مسجلة حتى الآن."
        if is_bot_msg:
            await target.edit_text(text)
        else:
            await target.answer(text, reply_markup=get_main_menu())
        return
    offset = page * EDIT_PAGE_SIZE
    invoices = await database.get_user_invoices_paginated(user_id, offset=offset, limit=EDIT_PAGE_SIZE)
    buttons = []
    for inv_id, created_at, amount, currency, category in invoices:
        date_str = str(created_at)[:10]
        cat = category or "عام"
        emoji = get_category_emoji(cat)
        label = f"{emoji} #{inv_id} | {date_str} | {amount:,.0f} {currency} | {cat}"
        buttons.append([InlineKeyboardButton(text=label, callback_data=f"inv_edit_{inv_id}")])
    nav_row = []
    if page > 0:
        nav_row.append(InlineKeyboardButton(text="⬅️ السابق", callback_data=f"editpage_{page-1}"))
    if offset + EDIT_PAGE_SIZE < total:
        nav_row.append(InlineKeyboardButton(text="التالي ➡️", callback_data=f"editpage_{page+1}"))
    if nav_row:
        buttons.append(nav_row)
    kb = InlineKeyboardMarkup(inline_keyboard=buttons)
    header = f"📋 فواتيرك ({total} فاتورة) — اختر للتعديل:"
    if is_bot_msg:
        await target.edit_text(header, reply_markup=kb)
    else:
        await target.answer(header, reply_markup=kb)

@dp.callback_query(F.data.startswith("editpage_"))
async def handle_edit_page(callback: CallbackQuery, state: FSMContext):
    page = int(callback.data.split("_")[1])
    await _show_invoice_list_for_edit(callback.message, callback.from_user.id, page=page)
    await callback.answer()

@dp.callback_query(F.data.startswith("inv_edit_"))
async def handle_edit_invoice_selected(callback: CallbackQuery, state: FSMContext):
    inv_id = int(callback.data.replace("inv_edit_", ""))
    invoice = await database.get_invoice_by_id(inv_id, callback.from_user.id)
    if not invoice:
        await callback.answer("الفاتورة دي مش موجودة!", show_alert=True)
        return
    _, created_at, amount, currency, category, raw_text, image_file_id = invoice
    date_str = str(created_at)[:10]
    cat = category or "عام"
    img_status = "✅ مرفقة" if image_file_id else "❌ لا توجد"
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="✏️ تعديل المبلغ",    callback_data=f"inv_edt_price_{inv_id}")],
        [InlineKeyboardButton(text="📝 تعديل التفاصيل",  callback_data=f"inv_edt_text_{inv_id}")],
        [InlineKeyboardButton(text="🗑️ حذف الفاتورة",   callback_data=f"inv_del_{inv_id}")],
        [InlineKeyboardButton(text="🔙 رجوع",            callback_data="editpage_0")],
    ])
    await callback.message.edit_text(
        f"📄 **تفاصيل الفاتورة #{inv_id}**\n\n"
        f"📅 التاريخ: {date_str}\n"
        f"🏷️ القسم: {cat}\n"
        f"💰 المبلغ: {amount:,.2f} {currency}\n"
        f"📝 التفاصيل: {raw_text or 'لا يوجد'}\n"
        f"📸 الصورة: {img_status}\n\n"
        "اختر ما تريد:",
        reply_markup=kb,
        parse_mode='Markdown'
    )
    await callback.answer()

@dp.callback_query(F.data.startswith("inv_edt_price_") | F.data.startswith("inv_edt_text_"))
async def handle_edit_field_request(callback: CallbackQuery, state: FSMContext):
    parts = callback.data.split("_")
    # inv_edt_price_<id>  or  inv_edt_text_<id>
    field = parts[2]   # "price" or "text"
    inv_id = int(parts[3])
    await state.set_state(MarineStates.waiting_for_edit_value)
    await state.update_data(edit_inv_id=inv_id, edit_field=field)
    prompt = "💰 اكتب المبلغ الجديد:" if field == "price" else "📝 اكتب التفاصيل الجديدة:"
    await callback.message.edit_text(prompt)
    await callback.answer()

@dp.message(MarineStates.waiting_for_edit_value, F.text & ~F.text.startswith('/'))
async def handle_edit_value_input(message: Message, state: FSMContext):
    data = await state.get_data()
    inv_id = data.get('edit_inv_id')
    field  = data.get('edit_field')
    text   = message.text.strip()
    if field == 'price':
        amount = _extract_amount(text)
        await database.update_invoice(inv_id, message.from_user.id, amount=amount)
        reply = f"✅ تم تحديث المبلغ إلى {amount:,.2f} جنيه للفاتورة #{inv_id}"
    else:
        await database.update_invoice(inv_id, message.from_user.id, raw_text=text)
        reply = f"✅ تم تحديث تفاصيل الفاتورة #{inv_id}"
    await state.clear()
    await message.answer(reply, reply_markup=get_main_menu())

@dp.callback_query(F.data.startswith("inv_del_") & ~F.data.startswith("inv_del_confirm_"))
async def handle_delete_single_invoice(callback: CallbackQuery, state: FSMContext):
    inv_id = int(callback.data.replace("inv_del_", ""))
    kb = InlineKeyboardMarkup(inline_keyboard=[[
        InlineKeyboardButton(text="✅ نعم، احذف",  callback_data=f"inv_del_confirm_{inv_id}"),
        InlineKeyboardButton(text="❌ إلغاء",       callback_data="editpage_0"),
    ]])
    await callback.message.edit_text(
        f"⚠️ هل أنت متأكد من حذف الفاتورة #{inv_id}؟", reply_markup=kb
    )
    await callback.answer()

@dp.callback_query(F.data.startswith("inv_del_confirm_"))
async def handle_delete_invoice_confirmed(callback: CallbackQuery, state: FSMContext):
    inv_id = int(callback.data.replace("inv_del_confirm_", ""))
    success = await database.delete_invoice_by_id(inv_id, callback.from_user.id)
    msg = f"✅ تم حذف الفاتورة #{inv_id}." if success else "❌ لم يتم العثور على الفاتورة."
    await callback.message.edit_text(msg)
    await callback.answer()

# ─────────────────────────────────────────────────────────────
# REQ #4 — مسح  (Delete Management)
# ─────────────────────────────────────────────────────────────
@dp.message(F.text == "🗑️ مسح")
async def handle_delete_menu(message: Message, state: FSMContext):
    await state.clear()
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="🔥 مسح كل الفواتير",          callback_data="del_all")],
        [InlineKeyboardButton(text="📅 مسح فواتير شهر معين",      callback_data="del_month")],
        [InlineKeyboardButton(text="📋 اختيار فاتورة للحذف",      callback_data="del_single")],
    ])
    await message.answer("🗑️ **إدارة الحذف** — اختر:", reply_markup=kb, parse_mode='Markdown')

# — Delete All ─────────────────────────────────────────────────
@dp.callback_query(F.data == "del_all")
async def del_all_confirm_prompt(callback: CallbackQuery, state: FSMContext):
    kb = InlineKeyboardMarkup(inline_keyboard=[[
        InlineKeyboardButton(text="✅ نعم، امسح الكل", callback_data="del_all_confirm"),
        InlineKeyboardButton(text="❌ إلغاء",           callback_data="del_cancel"),
    ]])
    await callback.message.edit_text(
        "⚠️ **تحذير!** هل أنت متأكد من مسح **جميع** فواتيرك؟\nهذا الإجراء لا يمكن التراجع عنه.",
        reply_markup=kb, parse_mode='Markdown'
    )
    await callback.answer()

@dp.callback_query(F.data == "del_all_confirm")
async def del_all_execute(callback: CallbackQuery, state: FSMContext):
    count = await database.delete_all_invoices(callback.from_user.id)
    await callback.message.edit_text(f"✅ تم مسح {count} فاتورة بنجاح.")
    await callback.answer()

# — Delete by Month ────────────────────────────────────────────
@dp.callback_query(F.data == "del_month")
async def del_month_picker(callback: CallbackQuery, state: FSMContext):
    now = datetime.now()
    buttons = []
    row = []
    MONTH_AR = ["يناير","فبراير","مارس","أبريل","مايو","يونيو",
                "يوليو","أغسطس","سبتمبر","أكتوبر","نوفمبر","ديسمبر"]
    for i in range(12):
        month_date = now - timedelta(days=30 * i)
        y, m = month_date.year, month_date.month
        label = f"{MONTH_AR[m-1]} {y}"
        row.append(InlineKeyboardButton(text=label, callback_data=f"del_month_{y}_{m}"))
        if len(row) == 2:
            buttons.append(row)
            row = []
    if row:
        buttons.append(row)
    buttons.append([InlineKeyboardButton(text="❌ إلغاء", callback_data="del_cancel")])
    await callback.message.edit_text(
        "📅 اختر الشهر الذي تريد مسح فواتيره:",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons)
    )
    await callback.answer()

@dp.callback_query(F.data.startswith("del_month_"))
async def del_month_execute(callback: CallbackQuery, state: FSMContext):
    _, _, year_str, month_str = callback.data.split("_")
    year, month = int(year_str), int(month_str)
    count = await database.delete_invoices_by_month(callback.from_user.id, year, month)
    MONTH_AR = ["يناير","فبراير","مارس","أبريل","مايو","يونيو",
                "يوليو","أغسطس","سبتمبر","أكتوبر","نوفمبر","ديسمبر"]
    month_name = MONTH_AR[month - 1]
    await callback.message.edit_text(f"✅ تم مسح {count} فاتورة من شهر {month_name} {year}.")
    await callback.answer()

# — Delete Single (reuse edit list) ───────────────────────────
@dp.callback_query(F.data == "del_single")
async def del_single_list(callback: CallbackQuery, state: FSMContext):
    user_id = callback.from_user.id
    total = await database.count_user_invoices(user_id)
    if total == 0:
        await callback.message.edit_text("لا توجد فواتير للحذف.")
        await callback.answer()
        return
    invoices = await database.get_user_invoices_paginated(user_id, offset=0, limit=EDIT_PAGE_SIZE)
    buttons = []
    for inv_id, created_at, amount, currency, category in invoices:
        date_str = str(created_at)[:10]
        cat = category or "عام"
        emoji = get_category_emoji(cat)
        label = f"{emoji} #{inv_id} | {date_str} | {amount:,.0f} {currency}"
        buttons.append([InlineKeyboardButton(text=label, callback_data=f"inv_del_{inv_id}")])
    buttons.append([InlineKeyboardButton(text="❌ إلغاء", callback_data="del_cancel")])
    await callback.message.edit_text(
        "📋 اختر الفاتورة للحذف:",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons)
    )
    await callback.answer()

@dp.callback_query(F.data == "del_cancel")
async def del_cancel(callback: CallbackQuery, state: FSMContext):
    await callback.message.edit_text("❌ تم إلغاء العملية.")
    await callback.answer()

# ─────────────────────────────────────────────────────────────
# REQ #5 — إضافة صورة الفاتورة  (Invoice Image) — v4.0 Paginated
# ─────────────────────────────────────────────────────────────
IMG_PAGE_SIZE = 5

@dp.message(F.text == "📸 إضافة صورة الفاتورة")
async def handle_add_image_menu(message: Message, state: FSMContext):
    await state.clear()
    await state.set_state(MarineStates.waiting_for_image_invoice_sel)
    await _show_invoice_list_for_image(message, message.from_user.id, page=0)

async def _show_invoice_list_for_image(target, user_id: int, page: int = 0):
    """Paginated invoice selector for attaching images (5 per page)."""
    total = await database.count_user_invoices(user_id)
    is_bot_msg = hasattr(target, 'from_user') and target.from_user and target.from_user.is_bot
    if total == 0:
        text = "لا توجد فواتير مسجلة. أضف فاتورة أولاً."
        if is_bot_msg:
            await target.edit_text(text)
        else:
            await target.answer(text, reply_markup=get_main_menu())
        return
    offset = page * IMG_PAGE_SIZE
    invoices = await database.get_user_invoices_paginated(user_id, offset=offset, limit=IMG_PAGE_SIZE)
    buttons = []
    for inv_id, created_at, amount, currency, category in invoices:
        date_str = str(created_at)[:10]
        cat = category or "عام"
        emoji = get_category_emoji(cat)
        label = f"{emoji} #{inv_id} | {date_str} | {amount:,.0f} | {cat}"
        buttons.append([InlineKeyboardButton(text=label, callback_data=f"img_inv_{inv_id}")])
    # Navigation row
    nav_row = []
    if page > 0:
        nav_row.append(InlineKeyboardButton(text="⬅️ السابق", callback_data=f"imgpage_{page-1}"))
    if offset + IMG_PAGE_SIZE < total:
        nav_row.append(InlineKeyboardButton(text="التالي ➡️", callback_data=f"imgpage_{page+1}"))
    if nav_row:
        buttons.append(nav_row)
    kb = InlineKeyboardMarkup(inline_keyboard=buttons)
    header = f"📸 اختر الفاتورة لإرفاق صورة ({total} فاتورة):"
    if is_bot_msg:
        await target.edit_text(header, reply_markup=kb)
    else:
        await target.answer(header, reply_markup=kb)

@dp.callback_query(F.data.startswith("imgpage_"))
async def handle_img_page(callback: CallbackQuery, state: FSMContext):
    """Handle pagination for invoice image attachment list."""
    page = int(callback.data.split("_")[1])
    await _show_invoice_list_for_image(callback.message, callback.from_user.id, page=page)
    await callback.answer()

@dp.callback_query(MarineStates.waiting_for_image_invoice_sel, F.data.startswith("img_inv_"))
async def handle_image_invoice_selected(callback: CallbackQuery, state: FSMContext):
    inv_id = int(callback.data.replace("img_inv_", ""))
    invoice = await database.get_invoice_by_id(inv_id, callback.from_user.id)
    if not invoice:
        await callback.answer("الفاتورة مش موجودة!", show_alert=True)
        return
    await state.update_data(target_invoice_id=inv_id)
    await state.set_state(MarineStates.waiting_for_image_upload)
    await callback.message.edit_text(
        f"📸 أرسل صورة الفاتورة #{inv_id} الآن وسأقوم بربطها بها."
    )
    await callback.answer()

@dp.message(MarineStates.waiting_for_image_upload, F.photo)
async def handle_invoice_image_upload(message: Message, state: FSMContext):
    data = await state.get_data()
    inv_id = data.get('target_invoice_id')
    if not inv_id:
        await state.clear()
        return
    file_id = message.photo[-1].file_id
    await database.update_invoice_image(inv_id, message.from_user.id, file_id)
    await state.clear()
    await message.answer(
        f"✅ تم ربط الصورة بالفاتورة #{inv_id} بنجاح! 📎",
        reply_markup=get_main_menu()
    )

# ─────────────────────────────────────────────────────────────
# UTILITIES
# ─────────────────────────────────────────────────────────────
def _extract_amount(text: str) -> float:
    """Extracts the first numeric value from a string."""
    match = re.search(r'([\d,]+(?:[.,]\d+)?)', text)
    if match:
        try:
            return float(match.group(1).replace(',', ''))
        except ValueError:
            pass
    return 0.0

# ─────────────────────────────────────────────────────────────
# AI ENGINE
# ─────────────────────────────────────────────────────────────
def analyze_with_ai(input_text: str, mode: str = 'extract', boat_name: str = "مركب غير مسمى"):
    """Uses Gemini to extract invoice data or chat. Returns dict or string."""
    if mode == 'text_intent':
        text_lower = input_text.lower()
        num = _extract_amount(input_text)
        if 'بنزين' in text_lower or 'سولار' in text_lower or 'جاز' in text_lower:
            return {"action": "expense", "amount": num, "category": "بنزين"}
        if 'ماركت' in text_lower or 'أكل' in text_lower or 'شرب' in text_lower:
            return {"action": "expense", "amount": num, "category": "ماركت"}

    try:
        gemini_key = os.getenv("GEMINI_API_KEY")
        if not gemini_key:
            raise ValueError("GEMINI_API_KEY missing")

        if mode in ['extract', 'retry_extract', 'total_only', 'full_storage', 'text_intent']:
            if mode == 'text_intent':
                prompt = (f"Extract price and category from this Arabic text. "
                          f"Categories: [بنزين,صيانة,ماركت,إكرامية,أدوات نظافة,عام]. "
                          f"Return ONLY JSON: {{\"action\":\"expense\",\"amount\":float,\"category\":\"string\"}}. "
                          f"Text: {input_text}")
            else:
                extra = "\\nWARNING: Look closer, find every item.\\n" if mode == 'retry_extract' else ""
                mode_desc = ("Return JSON with store, date, total only." if mode == 'total_only'
                             else "Return full JSON with all items, prices, total.")
                prompt = (f"{extra}You are an Arabic invoice OCR analyst. {mode_desc}\\n"
                          f'JSON format: {{"mode":"total/full","store":"str","date":"YYYY-MM-DD",'
                          f'"total":float,"items":[{{"name":"str","price":float,"qty":int}}]}}\\n'
                          f"Return ONLY valid JSON. Input:\\n{input_text}")
            system = f"أنت محاسب AbsyCode لمركب '{boat_name}'. أجب بـ JSON فقط."
        else:
            prompt = input_text
            system = (f"You are AbsyCode Assistant for boat '{boat_name}'. "
                      f"Speak in friendly Egyptian Arabic. Be concise.")

        url = f"https://generativelanguage.googleapis.com/v1/models/gemini-2.5-flash:generateContent?key={gemini_key}"
        payload = {
            "contents": [{"parts": [{"text": f"{system}\\n\\n{prompt}"}]}]
        }
        
        response_text = None
        resp = requests.post(url, json=payload, timeout=15)
        if resp.status_code == 200:
            data = resp.json()
            if "candidates" in data and len(data["candidates"]) > 0:
                response_text = data["candidates"][0]["content"]["parts"][0]["text"].strip()
        else:
            logging.warning(f"Gemini API returned {resp.status_code}: {resp.text}")

        if not response_text:
            return None
            
        content = response_text

    except Exception as e:
        logging.error(f"AI error ({mode}): {e}")
        if mode == 'text_intent':
            num = _extract_amount(input_text)
            return {"action": "add", "amount": num, "category": "عام"}
        if mode in ['extract', 'retry_extract', 'total_only', 'full_storage']:
            return {"items": [], "total": 0.0}
        return "حصل خطأ بسيط، حاول تاني. 😅"

# ─────────────────────────────────────────────────────────────
# /report  — v4.0 Detailed Ledger + CRUD Editing + Merge
# ─────────────────────────────────────────────────────────────
REPORT_PAGE_SIZE = 5

@dp.message(Command("report"))
@dp.message(F.text == "📊 التقارير")
async def cmd_report(message: Message):
    await _show_report_page(message, message.from_user.id, page=0)

async def _build_report_text(user_id: int) -> str:
    """Build the category summary header text (HTML)."""
    import html as html_mod
    user_data = await database.get_user(user_id)
    boat_name = user_data[2] if user_data and user_data[2] else "مركب غير مسمى"
    report_data = await database.get_detailed_report(user_id)
    total_invoices = await database.count_user_invoices(user_id)
    if not report_data or (len(report_data) == 1 and report_data[0][1] is None):
        return f"لا توجد مصاريف مسجلة لمركب {html_mod.escape(boat_name)} حتى الآن.", 0
    msg = f"📊 <b>تقرير {html_mod.escape(boat_name)}:</b>\n\n"
    totals = {}
    for category, amount, currency in report_data:
        if amount is None:
            continue
        curr = currency or "جنيه"
        cat = category or "أخرى"
        emoji = get_category_emoji(category)
        msg += f"{emoji} <b>{html_mod.escape(cat)}</b>: {amount:,.2f} {html_mod.escape(curr)}\n"
        totals[curr] = totals.get(curr, 0) + amount
    msg += "\n" + "━"*18 + "\n💰 <b>الإجمالي:</b> "
    parts = [f"{total:,.2f} {html_mod.escape(curr)}" for curr, total in totals.items()]
    msg += " | ".join(parts)
    msg += f"\n📋 <b>{total_invoices} فاتورة</b> — اضغط للتعديل:\n"
    return msg, total_invoices

async def _show_report_page(target, user_id: int, page: int = 0):
    """Shows category summary + paginated ledger entries with edit buttons."""
    header, total = await _build_report_text(user_id)
    is_bot_msg = hasattr(target, 'from_user') and target.from_user and target.from_user.is_bot
    if total == 0:
        if is_bot_msg:
            await target.edit_text(header, parse_mode='HTML')
        else:
            await target.answer(header, parse_mode='HTML')
        return
    # Fetch paginated ledger entries (newest first)
    offset = page * REPORT_PAGE_SIZE
    entries = await database.get_ledger_entries_paginated(user_id, offset=offset, limit=REPORT_PAGE_SIZE)
    buttons = []
    for inv_id, created_at, amount, currency, category, raw_text, image_file_id in entries:
        date_str = str(created_at)[:10]
        cat = category or "عام"
        emoji = get_category_emoji(cat)
        img_icon = "📸" if image_file_id else ""
        label = f"{emoji} #{inv_id} | {date_str} | {amount:,.0f} {currency or 'ج'} {img_icon}"
        buttons.append([
            InlineKeyboardButton(text=label, callback_data=f"rpt_edit_{inv_id}"),
        ])
    # Navigation row
    nav_row = []
    if page > 0:
        nav_row.append(InlineKeyboardButton(text="⬅️ السابق", callback_data=f"rptpage_{page-1}"))
    if offset + REPORT_PAGE_SIZE < total:
        nav_row.append(InlineKeyboardButton(text="التالي ➡️", callback_data=f"rptpage_{page+1}"))
    if nav_row:
        buttons.append(nav_row)
    kb = InlineKeyboardMarkup(inline_keyboard=buttons)
    if is_bot_msg:
        await target.edit_text(header, reply_markup=kb, parse_mode='HTML')
    else:
        await target.answer(header, reply_markup=kb, parse_mode='HTML')

@dp.callback_query(F.data.startswith("rptpage_"))
async def handle_report_page(callback: CallbackQuery):
    """Navigate report ledger pages."""
    page = int(callback.data.split("_")[1])
    await _show_report_page(callback.message, callback.from_user.id, page=page)
    await callback.answer()

# ── Report: Invoice Detail & Edit Sub-menu ────────────────────
@dp.callback_query(F.data.startswith("rpt_edit_"))
async def handle_report_edit_invoice(callback: CallbackQuery, state: FSMContext):
    """Show invoice detail with edit/merge options from report view."""
    inv_id = int(callback.data.replace("rpt_edit_", ""))
    invoice = await database.get_invoice_by_id(inv_id, callback.from_user.id)
    if not invoice:
        await callback.answer("الفاتورة مش موجودة!", show_alert=True)
        return
    _, created_at, amount, currency, category, raw_text, image_file_id = invoice
    date_str = str(created_at)[:10]
    cat = category or "عام"
    img_status = "✅ مرفقة" if image_file_id else "❌ لا توجد"
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="✏️ تعديل المبلغ",    callback_data=f"rpt_edt_price_{inv_id}")],
        [InlineKeyboardButton(text="📝 تعديل التفاصيل",  callback_data=f"rpt_edt_text_{inv_id}")],
        [InlineKeyboardButton(text="🔗 دمج مع فاتورة أخرى", callback_data=f"rpt_merge_{inv_id}")],
        [InlineKeyboardButton(text="🗑️ حذف الفاتورة",   callback_data=f"rpt_del_{inv_id}")],
        [InlineKeyboardButton(text="🔙 رجوع للتقرير",    callback_data="rptpage_0")],
    ])
    await callback.message.edit_text(
        f"📄 <b>تفاصيل الفاتورة #{inv_id}</b>\n\n"
        f"📅 التاريخ: {date_str}\n"
        f"🏷️ القسم: {cat}\n"
        f"💰 المبلغ: {amount:,.2f} {currency or 'جنيه'}\n"
        f"📝 التفاصيل: {raw_text or 'لا يوجد'}\n"
        f"📸 الصورة: {img_status}\n\n"
        "اختر ما تريد:",
        reply_markup=kb, parse_mode='HTML'
    )
    await callback.answer()

# ── Report: Edit Price / Text ─────────────────────────────────
@dp.callback_query(F.data.startswith("rpt_edt_price_") | F.data.startswith("rpt_edt_text_"))
async def handle_report_edit_field(callback: CallbackQuery, state: FSMContext):
    """Start editing a field from the report detail view."""
    parts = callback.data.split("_")
    field = parts[2]   # "price" or "text"
    inv_id = int(parts[3])
    await state.set_state(MarineStates.waiting_for_report_edit_value)
    await state.update_data(rpt_edit_inv_id=inv_id, rpt_edit_field=field)
    prompt = "💰 اكتب المبلغ الجديد:" if field == "price" else "📝 اكتب التفاصيل الجديدة:"
    await callback.message.edit_text(prompt)
    await callback.answer()

@dp.message(MarineStates.waiting_for_report_edit_value, F.text & ~F.text.startswith('/'))
async def handle_report_edit_value_input(message: Message, state: FSMContext):
    """Process the new value for report-based invoice editing."""
    data = await state.get_data()
    inv_id = data.get('rpt_edit_inv_id')
    field  = data.get('rpt_edit_field')
    text   = message.text.strip()
    if field == 'price':
        amount = _extract_amount(text)
        await database.update_invoice(inv_id, message.from_user.id, amount=amount)
        reply = f"✅ تم تحديث المبلغ إلى {amount:,.2f} جنيه للفاتورة #{inv_id}"
    else:
        await database.update_invoice(inv_id, message.from_user.id, raw_text=text)
        reply = f"✅ تم تحديث تفاصيل الفاتورة #{inv_id}"
    await state.clear()
    await message.answer(reply, reply_markup=get_main_menu())

# ── Report: Delete from report ────────────────────────────────
@dp.callback_query(F.data.startswith("rpt_del_") & ~F.data.startswith("rpt_del_confirm_"))
async def handle_report_delete_prompt(callback: CallbackQuery):
    inv_id = int(callback.data.replace("rpt_del_", ""))
    kb = InlineKeyboardMarkup(inline_keyboard=[[
        InlineKeyboardButton(text="✅ نعم، احذف",  callback_data=f"rpt_del_confirm_{inv_id}"),
        InlineKeyboardButton(text="❌ إلغاء",       callback_data="rptpage_0"),
    ]])
    await callback.message.edit_text(
        f"⚠️ هل أنت متأكد من حذف الفاتورة #{inv_id}؟", reply_markup=kb
    )
    await callback.answer()

@dp.callback_query(F.data.startswith("rpt_del_confirm_"))
async def handle_report_delete_confirmed(callback: CallbackQuery):
    inv_id = int(callback.data.replace("rpt_del_confirm_", ""))
    success = await database.delete_invoice_by_id(inv_id, callback.from_user.id)
    msg = f"✅ تم حذف الفاتورة #{inv_id}." if success else "❌ لم يتم العثور على الفاتورة."
    await callback.message.edit_text(msg)
    await callback.answer()

# ── Report: Merge Invoices Flow ───────────────────────────────
MERGE_PAGE_SIZE = 5

@dp.callback_query(F.data.startswith("rpt_merge_") & ~F.data.startswith("rpt_merge_sel_") & ~F.data.startswith("rpt_merge_confirm_") & ~F.data.startswith("rpt_merge_pg_"))
async def handle_merge_start(callback: CallbackQuery, state: FSMContext):
    """Start the merge flow — show list of other invoices to merge into."""
    source_id = int(callback.data.replace("rpt_merge_", ""))
    await state.set_state(MarineStates.waiting_for_merge_target)
    await state.update_data(merge_source_id=source_id)
    await _show_merge_target_list(callback.message, callback.from_user.id, source_id, page=0)
    await callback.answer()

async def _show_merge_target_list(target, user_id: int, source_id: int, page: int = 0):
    """Show paginated list of invoices to merge the source into (excluding source)."""
    total = await database.count_user_invoices(user_id)
    total_targets = total - 1  # exclude source
    if total_targets <= 0:
        await target.edit_text("لا توجد فواتير أخرى للدمج معها.")
        return
    offset = page * MERGE_PAGE_SIZE
    all_inv = await database.get_user_invoices_paginated(user_id, offset=0, limit=200)
    # Filter out the source invoice and apply manual pagination
    targets = [(i, c, a, cu, cat) for i, c, a, cu, cat in all_inv if i != source_id]
    page_targets = targets[offset:offset + MERGE_PAGE_SIZE]
    buttons = []
    for inv_id, created_at, amount, currency, category in page_targets:
        date_str = str(created_at)[:10]
        cat = category or "عام"
        emoji = get_category_emoji(cat)
        label = f"{emoji} #{inv_id} | {date_str} | {amount:,.0f} {currency or 'ج'}"
        buttons.append([InlineKeyboardButton(
            text=label, callback_data=f"rpt_merge_sel_{source_id}_{inv_id}"
        )])
    nav_row = []
    if page > 0:
        nav_row.append(InlineKeyboardButton(text="⬅️ السابق", callback_data=f"rpt_merge_pg_{source_id}_{page-1}"))
    if offset + MERGE_PAGE_SIZE < len(targets):
        nav_row.append(InlineKeyboardButton(text="التالي ➡️", callback_data=f"rpt_merge_pg_{source_id}_{page+1}"))
    if nav_row:
        buttons.append(nav_row)
    buttons.append([InlineKeyboardButton(text="❌ إلغاء", callback_data="rptpage_0")])
    kb = InlineKeyboardMarkup(inline_keyboard=buttons)
    await target.edit_text(
        f"🔗 <b>دمج الفاتورة #{source_id}</b>\n\n"
        "اختر الفاتورة التي تريد الدمج فيها:",
        reply_markup=kb, parse_mode='HTML'
    )

@dp.callback_query(F.data.startswith("rpt_merge_pg_"))
async def handle_merge_page(callback: CallbackQuery, state: FSMContext):
    """Navigate merge target list pages."""
    # rpt_merge_pg_{source_id}_{page}
    parts = callback.data.split("_")
    source_id = int(parts[3])
    page = int(parts[4])
    await _show_merge_target_list(callback.message, callback.from_user.id, source_id, page=page)
    await callback.answer()

@dp.callback_query(F.data.startswith("rpt_merge_sel_"))
async def handle_merge_target_selected(callback: CallbackQuery, state: FSMContext):
    """User selected a merge target — show confirmation."""
    # rpt_merge_sel_{source_id}_{target_id}
    parts = callback.data.split("_")
    source_id = int(parts[3])
    target_id = int(parts[4])
    source = await database.get_invoice_by_id(source_id, callback.from_user.id)
    target = await database.get_invoice_by_id(target_id, callback.from_user.id)
    if not source or not target:
        await callback.answer("فاتورة غير موجودة!", show_alert=True)
        return
    kb = InlineKeyboardMarkup(inline_keyboard=[[
        InlineKeyboardButton(text="✅ نعم، ادمج", callback_data=f"rpt_merge_confirm_{source_id}_{target_id}"),
        InlineKeyboardButton(text="❌ إلغاء",     callback_data="rptpage_0"),
    ]])
    new_total = (source[2] or 0) + (target[2] or 0)
    await callback.message.edit_text(
        f"⚠️ <b>تأكيد الدمج</b>\n\n"
        f"📤 الفاتورة #{source_id} ({source[2]:,.2f} {source[3] or 'جنيه'})\n"
        f"📥 ← تُدمج في → الفاتورة #{target_id} ({target[2]:,.2f} {target[3] or 'جنيه'})\n\n"
        f"💰 الإجمالي بعد الدمج: <b>{new_total:,.2f} {target[3] or 'جنيه'}</b>\n\n"
        "⚠️ سيتم حذف الفاتورة المصدر بعد الدمج.",
        reply_markup=kb, parse_mode='HTML'
    )
    await callback.answer()

@dp.callback_query(F.data.startswith("rpt_merge_confirm_"))
async def handle_merge_execute(callback: CallbackQuery, state: FSMContext):
    """Execute the merge operation."""
    # rpt_merge_confirm_{source_id}_{target_id}
    parts = callback.data.split("_")
    source_id = int(parts[3])
    target_id = int(parts[4])
    result = await database.merge_invoices(callback.from_user.id, source_id, target_id)
    await state.clear()
    if result:
        await callback.message.edit_text(
            f"✅ <b>تم الدمج بنجاح!</b>\n\n"
            f"🔗 الفاتورة #{source_id} دُمجت في #{target_id}\n"
            f"💰 الإجمالي الجديد: {result['new_amount']:,.2f} {result['currency']}",
            parse_mode='HTML'
        )
    else:
        await callback.message.edit_text("❌ حدث خطأ أثناء الدمج. تأكد أن الفواتير موجودة.")
    await callback.answer()

# ─────────────────────────────────────────────────────────────
# /export  (with embedded invoice images)
# ─────────────────────────────────────────────────────────────
@dp.message(Command("export"))
async def cmd_export(message: Message):
    user_id = message.from_user.id
    proc = await message.answer("⏳ جاري تجهيز تقارير Excel و Word...")
    raw_invoices = await database.get_user_invoices(user_id)
    if not raw_invoices:
        await proc.edit_text("لا توجد مصاريف للتصدير.")
        return
    user_data = await database.get_user(user_id)
    boat_name = user_data[2] if user_data and user_data[2] else "مركب غير مسمى"
    excel_path = f"temp/Invoices_{user_id}.xlsx"
    word_path  = f"temp/Report_{user_id}.docx"
    img_paths  = []
    try:
        # ── Excel ──────────────────────────────────────────────
        df = pd.DataFrame(raw_invoices,
                          columns=['Date','Amount','Currency','Category','RawText','ImageFileId'])
        raw_items = await database.get_all_invoice_items(user_id)
        df_items  = pd.DataFrame(raw_items,
                                 columns=['Date','InvoiceID','Category','ItemName','Qty','Price'])
        with pd.ExcelWriter(excel_path, engine='openpyxl') as writer:
            df.drop(columns=['ImageFileId']).to_excel(writer, sheet_name='Invoices', index=False)
            df_items.to_excel(writer, sheet_name='Items_Detail', index=False)

        # Embed images into Excel
        try:
            from openpyxl import load_workbook
            from openpyxl.drawing.image import Image as XLImage
            invoices_with_imgs = await database.get_invoices_with_images(user_id)
            if invoices_with_imgs:
                wb = load_workbook(excel_path)
                ws = wb.create_sheet("Invoice Images")
                ws['A1'] = "رقم الفاتورة"
                ws['B1'] = "التاريخ"
                ws['C1'] = "المبلغ"
                ws['D1'] = "الصورة"
                for row_num, (inv_id, created_at, amount, category, file_id) in enumerate(invoices_with_imgs, start=2):
                    ws.row_dimensions[row_num].height = 100
                    ws[f'A{row_num}'] = inv_id
                    ws[f'B{row_num}'] = str(created_at)[:10]
                    ws[f'C{row_num}'] = amount
                    img_local = f"temp/img_export_{inv_id}.jpg"
                    img_paths.append(img_local)
                    tg_file = await bot.get_file(file_id)
                    await bot.download_file(tg_file.file_path, img_local)
                    xl_img = XLImage(img_local)
                    xl_img.width  = 120
                    xl_img.height = 90
                    ws.add_image(xl_img, f'D{row_num}')
                wb.save(excel_path)
        except Exception as img_e:
            logging.warning(f"Excel image embed failed: {img_e}")

        # ── Word ───────────────────────────────────────────────
        report_data = await database.get_detailed_report(user_id)
        doc = docx.Document()
        title = doc.add_heading(f'{boat_name} — AbsyCode Invoice Report', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph()
        if raw_items:
            h = doc.add_heading('تفاصيل الأصناف', level=2)
            h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            tbl = doc.add_table(rows=1, cols=4)
            tbl.style = 'Table Grid'
            hdr = tbl.rows[0].cells
            hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = 'التاريخ','القسم','الصنف','السعر'
            for date_str, _, cat, item_name, qty, price in raw_items:
                row = tbl.add_row().cells
                row[0].text = str(date_str)[:10]
                row[1].text = cat or "أخرى"
                row[2].text = str(item_name)
                row[3].text = f"{price:,.2f}"
            doc.add_paragraph()
        h2 = doc.add_heading('ملخص الأقسام', level=2)
        h2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        tbl2 = doc.add_table(rows=1, cols=3)
        tbl2.style = 'Table Grid'
        hdr2 = tbl2.rows[0].cells
        hdr2[0].text, hdr2[1].text, hdr2[2].text = 'القسم','المبلغ','العملة'
        totals = {}
        for category, amount, currency in report_data:
            if amount is None:
                continue
            r = tbl2.add_row().cells
            r[0].text = category or "أخرى"
            r[1].text = f"{amount:,.2f}"
            curr = currency or "جنيه"
            r[2].text = curr
            totals[curr] = totals.get(curr, 0) + amount
        doc.add_paragraph()
        for curr, total in totals.items():
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run = p.add_run(f"الإجمالي الكلي ({curr}): {total:,.2f}")
            run.bold = True
            run.font.size = Pt(14)

        # Embed invoice images in Word
        invoices_with_imgs = await database.get_invoices_with_images(user_id)
        if invoices_with_imgs:
            doc.add_page_break()
            hi = doc.add_heading('صور الفواتير', level=2)
            hi.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for inv_id, created_at, amount, category, file_id in invoices_with_imgs:
                doc.add_paragraph(f"فاتورة #{inv_id} — {str(created_at)[:10]} — {amount:,.2f}")
                img_local = f"temp/img_word_{inv_id}.jpg"
                img_paths.append(img_local)
                if not os.path.exists(img_local):
                    tg_file = await bot.get_file(file_id)
                    await bot.download_file(tg_file.file_path, img_local)
                try:
                    doc.add_picture(img_local, width=Inches(4))
                except Exception:
                    pass
                doc.add_paragraph()
        doc.save(word_path)

        await proc.edit_text("✅ جاري الإرسال...")
        await message.answer_document(FSInputFile(excel_path), caption="📊 تقرير Excel")
        await message.answer_document(FSInputFile(word_path),  caption="📝 تقرير Word")

    except Exception as e:
        logging.error(f"Export error: {e}")
        traceback.print_exc()
        await message.answer("❌ حدث خطأ أثناء التصدير.")
    finally:
        for p in [excel_path, word_path] + img_paths:
            try:
                if os.path.exists(p):
                    os.remove(p)
            except Exception:
                pass

# ─────────────────────────────────────────────────────────────
# PHOTO HANDLER (OCR) — fixes NameError bug (local_path vs local_filename)
# ─────────────────────────────────────────────────────────────
@dp.message(F.photo, ~StateFilter(MarineStates.waiting_for_image_upload))
async def handle_photo(message: Message, state: FSMContext):
    user_id = message.from_user.id
    if user_id in pending_invoices:
        await message.answer("عندك فاتورة بتتعالج، خلصها الأول!")
        return
        
    proc = await message.answer("جاري قراءة الفاتورة وتحليلها... 🧐")
    local_path = f"temp/invoice_{user_id}_{message.message_id}.jpg"
    try:
        file = await bot.get_file(message.photo[-1].file_id)
        await bot.download_file(file.file_path, local_path)
        await proc.edit_text("جاري استخراج البيانات بالذكاء الاصطناعي... 🔍")
        
        import base64
        with open(local_path, "rb") as f:
            image_data = base64.b64encode(f.read()).decode("utf-8")
            
        gemini_key = os.getenv("GEMINI_API_KEY")
        if not gemini_key:
            await proc.edit_text("مفتاح Gemini API مش متفعل.")
            return

        user_data = await database.get_user(user_id)
        boat_name = user_data[2] if user_data and user_data[2] else "مركب غير مسمى"
        
        system = f"أنت محاسب AbsyCode لمركب '{boat_name}'. أجب بـ JSON فقط."
        prompt = (f"You are an expert Arabic/English OCR analyst. Please analyze this invoice image. "
                  f"Extract the full text (raw_text), total amount (total), date (date), and items if any. "
                  f"Map the invoice to one of these categories: [بنزين, صيانة, ماركت, إكرامية, أدوات نظافة, عام] based on strict keyword mapping (e.g. fuel -> بنزين). "
                  f'Return ONLY valid JSON format: {{"raw_text":"str","total":float,"date":"YYYY-MM-DD","category":"str","items":[{{"name":"str","price":float,"qty":int}}]}}')

        url = f"https://generativelanguage.googleapis.com/v1/models/gemini-2.5-flash:generateContent?key={gemini_key}"
        payload = {
            "contents": [{
                "parts": [
                    {"text": f"{system}\\n\\n{prompt}"},
                    {"inlineData": {"mimeType": "image/jpeg", "data": image_data}}
                ]
            }]
        }
        
        loop = asyncio.get_running_loop()
        resp = await loop.run_in_executor(None, lambda: requests.post(url, json=payload, timeout=20))
        
        if resp.status_code == 200:
            data = resp.json()
            if "candidates" in data and len(data["candidates"]) > 0:
                response_text = data["candidates"][0]["content"]["parts"][0]["text"].strip()
                # Clean markdown blocks if present
                if response_text.startswith("```json"):
                    response_text = response_text[7:]
                if response_text.startswith("```"):
                    response_text = response_text[3:]
                if response_text.endswith("```"):
                    response_text = response_text[:-3]
                
                ai_data = json.loads(response_text.strip())
                raw_text = ai_data.get("raw_text", "لم يتم التعرف على النص")
                manual_total = float(ai_data.get("total", 0.0))
                manual_date = ai_data.get("date", "غير محدد")
                category = ai_data.get("category", "عام")
                items = ai_data.get("items", [])
                
                pending_invoices[user_id] = {
                    'raw_text': raw_text, 'manual_total': manual_total,
                    'manual_date': manual_date, 'category': category,
                    'items': items
                }
                kb = InlineKeyboardMarkup(inline_keyboard=[
                    [InlineKeyboardButton(text="📂 حفظ الفاتورة", callback_data="process_ai_direct")]
                ])
                await proc.edit_text(
                    f"📥 تم قراءة الفاتورة بنجاح!\n💰 الإجمالي: {manual_total:,.2f} جنيه\n"
                    f"📅 التاريخ: {manual_date}\n🏷️ القسم: {category}\n📦 الأصناف: {len(items)}\n\nهل تريد الحفظ؟", reply_markup=kb)
            else:
                await proc.edit_text("لم أتمكن من استخراج البيانات، حاول بصورة أوضح.")
        else:
            logging.error(f"Gemini API Error: {resp.text}")
            await proc.edit_text("حدث خطأ في الاتصال بالذكاء الاصطناعي.")
            
    except json.JSONDecodeError:
        await proc.edit_text("حدث خطأ في فهم استجابة الذكاء الاصطناعي، يرجى المحاولة مرة أخرى.")
    except Exception as e:
        logging.error(f"Photo error: {e}")
        await proc.edit_text("حدث خطأ أثناء معالجة الصورة.")
    finally:
        if os.path.exists(local_path):
            os.remove(local_path)

@dp.callback_query(F.data == "process_ai_direct")
async def process_ai_direct_choice(callback: CallbackQuery):
    user_id = callback.from_user.id
    if user_id not in pending_invoices:
        await callback.answer("الفاتورة قديمة، ابعتها تاني.", show_alert=True)
        return
    data  = pending_invoices[user_id]
    await callback.message.edit_text("جاري الحفظ... 💾")
    try:
        best_amt  = data['manual_total']
        items     = data.get('items', [])
        category  = data.get('category', 'عام')
        await database.save_invoice(user_id, best_amt, "جنيه", data['raw_text'], items=items, category=category)
        del pending_invoices[user_id]
        await callback.message.edit_text(
            f"✅ تم حفظ الفاتورة بنجاح!\n💰 الإجمالي: {best_amt:,.2f} جنيه\n"
            f"🏷️ القسم: {category}\n📦 الأصناف: {len(items)}\nاستخدم /report لعرض التقرير.")
    except Exception as e:
        logging.error(f"Process invoice error: {e}")
        await callback.message.edit_text("حصل خطأ أثناء الحفظ، جرب تاني.")
    await callback.answer()

# ─────────────────────────────────────────────────────────────
# VOICE HANDLER — v4.0 Interactive Voice-to-Text Textarea & Manual Injection
# ─────────────────────────────────────────────────────────────
pending_voice_transcripts: dict = {}

def _build_voice_confirm_kb() -> InlineKeyboardMarkup:
    """3-button inline keyboard for voice transcript confirmation."""
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(
            text="✅ نعم، حقن في الفاتورة المانويل",
            callback_data="voice_inject"
        )],
        [InlineKeyboardButton(
            text="✏️ تعديل النص أولاً",
            callback_data="voice_edit_text"
        )],
        [InlineKeyboardButton(
            text="❌ إلغاء العملية",
            callback_data="voice_cancel"
        )],
    ])

@dp.message(F.voice)
async def handle_voice(message: Message, state: FSMContext):
    user_id = message.from_user.id

    if user_id in pending_invoices:
        await message.answer("عندك فاتورة بتتعالج، خلصها الأول!")
        return

    wait_msg = await message.answer("🎙️ جاري الاستماع للرسالة الصوتية...")
    ogg_path = f"temp/voice_{user_id}_{message.message_id}.ogg"
    wav_path = f"temp/voice_{user_id}_{message.message_id}.wav"

    try:
        # 1. Download voice file from Telegram
        voice_file = await bot.get_file(message.voice.file_id)
        await bot.download_file(voice_file.file_path, ogg_path)

        await wait_msg.edit_text("🔄 جاري تحويل الصوت للنص...")

        # 2. Convert OGG/Opus → WAV using pydub (requires ffmpeg)
        from pydub import AudioSegment
        loop = asyncio.get_running_loop()
        audio_segment = await loop.run_in_executor(
            None, lambda: AudioSegment.from_ogg(ogg_path)
        )
        await loop.run_in_executor(
            None, lambda: audio_segment.export(wav_path, format="wav")
        )

        # 3. Speech-to-text using Google's free web API (supports Arabic)
        import speech_recognition as sr
        recognizer = sr.Recognizer()
        with sr.AudioFile(wav_path) as source:
            audio_data = recognizer.record(source)

        transcript = await loop.run_in_executor(
            None,
            lambda: recognizer.recognize_google(audio_data, language="ar-EG")
        )

        if not transcript or not transcript.strip():
            await wait_msg.edit_text("مقدرتش أفهم الصوت، ممكن تحاول تاني بصوت أوضح؟ 🎤")
            return

        transcript = transcript.strip()

        # 4. Store transcript and present 3-option interactive keyboard
        pending_voice_transcripts[user_id] = transcript
        amount = _extract_amount(transcript)

        await wait_msg.edit_text(
            f"🎙️ أنا سمعت: `{transcript}`\n\n"
            f"💰 المبلغ المستخرج: {amount:,.2f} جنيه\n\n"
            "✍️ يمكنك تعديل النص أعلاه الآن إذا كان هناك خطأ، "
            "أو اضغط على الأزرار لتأكيد وحقن البيانات.",
            reply_markup=_build_voice_confirm_kb(),
            parse_mode='Markdown'
        )

    except Exception as e:
        # Handle speech_recognition specific errors
        err_name = type(e).__name__
        if err_name == "UnknownValueError":
            await wait_msg.edit_text("مقدرتش أفهم الصوت، ممكن تحاول تاني بصوت أوضح؟ 🎤")
        elif err_name == "RequestError":
            logging.error(f"Google STT service error: {e}")
            await wait_msg.edit_text("حدث خطأ في خدمة التعرف على الصوت. حاول تاني.")
        else:
            logging.error(f"Voice handler error: {e}")
            traceback.print_exc()
            await wait_msg.edit_text("حصل خطأ أثناء معالجة الرسالة الصوتية. ✍️")
    finally:
        for p in [ogg_path, wav_path]:
            if os.path.exists(p):
                try:
                    os.remove(p)
                except Exception:
                    pass


@dp.callback_query(F.data == "voice_inject")
async def handle_voice_inject(callback: CallbackQuery, state: FSMContext):
    """Inject transcript directly into the manual invoice FSM flow.
    Extracts amount via regex, routes to category selection, then
    handle_invoice_category_selected picks up voice_transcript from state.
    """
    user_id = callback.from_user.id
    transcript = pending_voice_transcripts.pop(user_id, None)

    if not transcript:
        await callback.answer("انتهت صلاحية النص، ابعت الصوت تاني.", show_alert=True)
        return

    amount = _extract_amount(transcript)

    # Route into the stable manual invoice FSM: category → auto-save
    await state.clear()
    await state.set_state(MarineStates.waiting_for_invoice_category)
    await state.update_data(voice_transcript=transcript, voice_amount=amount)

    await callback.message.edit_text(
        f"🎙️ **تم التعرف على النص:**\n\n"
        f"📝 \"{transcript}\"\n"
        f"💰 المبلغ المستخرج: {amount:,.2f} جنيه\n\n"
        f"📂 اختر قسم الفاتورة:",
        reply_markup=get_category_keyboard(),
        parse_mode='Markdown'
    )
    await callback.answer()


@dp.callback_query(F.data == "voice_edit_text")
async def handle_voice_edit_text(callback: CallbackQuery, state: FSMContext):
    """Put the bot in correction state — user types the corrected text."""
    user_id = callback.from_user.id
    transcript = pending_voice_transcripts.get(user_id)
    if not transcript:
        await callback.answer("انتهت صلاحية النص، ابعت الصوت تاني.", show_alert=True)
        return

    await state.set_state(MarineStates.waiting_for_voice_correction)
    await callback.message.edit_text(
        f"✏️ **النص الحالي:**\n`{transcript}`\n\n"
        "اكتب النص الصحيح الآن وسأعيد عرض الخيارات:",
        parse_mode='Markdown'
    )
    await callback.answer()


@dp.message(MarineStates.waiting_for_voice_correction, F.text & ~F.text.startswith('/'))
async def handle_voice_correction_input(message: Message, state: FSMContext):
    """User typed the corrected transcript. Store it and re-show the 3-button keyboard."""
    user_id = message.from_user.id
    corrected = message.text.strip()
    pending_voice_transcripts[user_id] = corrected
    amount = _extract_amount(corrected)

    await state.clear()  # Clear correction state, back to neutral
    await message.answer(
        f"🎙️ أنا سمعت: `{corrected}`\n\n"
        f"💰 المبلغ المستخرج: {amount:,.2f} جنيه\n\n"
        "✍️ يمكنك تعديل النص أعلاه الآن إذا كان هناك خطأ، "
        "أو اضغط على الأزرار لتأكيد وحقن البيانات.",
        reply_markup=_build_voice_confirm_kb(),
        parse_mode='Markdown'
    )


@dp.callback_query(F.data == "voice_cancel")
async def handle_voice_cancel(callback: CallbackQuery, state: FSMContext):
    """User cancelled the voice operation — clean up and dismiss."""
    pending_voice_transcripts.pop(callback.from_user.id, None)
    await state.clear()
    await callback.message.edit_text("❌ تم إلغاء العملية.")
    await callback.answer()


# ─────────────────────────────────────────────────────────────
# ADMIN PANEL (In-Bot Mobile Control)
# ─────────────────────────────────────────────────────────────
def _admin_main_kb() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="👥 عرض المستخدمين",             callback_data="adm_users")],
        [InlineKeyboardButton(text="⏳ تعديل الاشتراكات",           callback_data="adm_subs")],
        [InlineKeyboardButton(text="📜 سجل العمليات (Audit Log)",   callback_data="adm_audit")],
    ])

@dp.message(Command("admin"))
async def cmd_admin(message: Message, state: FSMContext):
    if message.from_user.id != SUPER_ADMIN_ID:
        return
    await state.clear()
    await message.answer(
        "🔐 <b>لوحة تحكم المشرف — AbsyCode</b>\n\nاختر من القائمة:",
        reply_markup=_admin_main_kb(), parse_mode='HTML'
    )

# ── View Users ────────────────────────────────────────────────
@dp.callback_query(F.data == "adm_users")
async def adm_view_users(callback: CallbackQuery):
    if callback.from_user.id != SUPER_ADMIN_ID:
        await callback.answer("غير مصرح.", show_alert=True); return
    import html
    users = await database.get_all_users_stats()
    if not users:
        await callback.message.edit_text("لا يوجد مستخدمون مسجلون.", reply_markup=_admin_main_kb())
        await callback.answer(); return
    now = datetime.now()
    msg = "👥 <b>قائمة المستخدمين:</b>\n\n"
    for uid, boat, is_act, expiry, reg_at in users:
        status = "❌ موقوف"
        if is_act and expiry:
            try:
                status = "✅ نشط" if datetime.fromisoformat(expiry) > now else "⏰ منتهي"
            except ValueError:
                status = "⚠️ خطأ"
        exp_str = str(expiry)[:10] if expiry else "—"
        inv_count = await database.count_user_invoices(uid)
        msg += (f"🆔 <code>{uid}</code>\n"
                f"🛥️ {html.escape(str(boat))} | {status}\n"
                f"📅 حتى {exp_str} | 📄 {inv_count} فاتورة\n\n")
    await callback.message.edit_text(msg, reply_markup=InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="🔙 رجوع", callback_data="adm_back")]
    ]), parse_mode='HTML')
    await callback.answer()

# ── Subscription Management ──────────────────────────────────
@dp.callback_query(F.data == "adm_subs")
async def adm_subs_list(callback: CallbackQuery):
    if callback.from_user.id != SUPER_ADMIN_ID:
        await callback.answer("غير مصرح.", show_alert=True); return
    users = await database.get_all_users_stats()
    if not users:
        await callback.message.edit_text("لا يوجد مستخدمون.", reply_markup=_admin_main_kb())
        await callback.answer(); return
    buttons = []
    for uid, boat, is_act, expiry, _ in users:
        status_icon = "✅" if is_act else "❌"
        label = f"{status_icon} {uid} — {boat}"
        buttons.append([InlineKeyboardButton(text=label, callback_data=f"adm_sub_{uid}")])
    buttons.append([InlineKeyboardButton(text="🔙 رجوع", callback_data="adm_back")])
    await callback.message.edit_text(
        "⏳ <b>تعديل الاشتراكات</b>\nاختر المستخدم:",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons), parse_mode='HTML'
    )
    await callback.answer()

@dp.callback_query(F.data.startswith("adm_sub_"))
async def adm_sub_user_menu(callback: CallbackQuery):
    if callback.from_user.id != SUPER_ADMIN_ID:
        await callback.answer("غير مصرح.", show_alert=True); return
    uid = int(callback.data.replace("adm_sub_", ""))
    user_data = await database.get_user(uid)
    if not user_data:
        await callback.answer("المستخدم غير موجود.", show_alert=True); return
    is_act, expiry, boat = user_data
    exp_str = str(expiry)[:10] if expiry else "—"
    now = datetime.now()
    status = "❌ موقوف"
    if is_act and expiry:
        try:
            status = "✅ نشط" if datetime.fromisoformat(expiry) > now else "⏰ منتهي"
        except ValueError:
            status = "⚠️ خطأ"
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="➕ +30 يوم",   callback_data=f"adm_ext30_{uid}"),
         InlineKeyboardButton(text="➖ -7 أيام",   callback_data=f"adm_red7_{uid}")],
        [InlineKeyboardButton(text="📅 تاريخ محدد", callback_data=f"adm_setexp_{uid}")],
        [InlineKeyboardButton(text="🚫 إيقاف",     callback_data=f"adm_revoke_{uid}")],
        [InlineKeyboardButton(text="🔙 رجوع",      callback_data="adm_subs")],
    ])
    await callback.message.edit_text(
        f"👤 <b>إدارة اشتراك</b>\n\n"
        f"🆔 <code>{uid}</code>\n"
        f"🛥️ {boat}\n"
        f"⏳ {status} حتى {exp_str}",
        reply_markup=kb, parse_mode='HTML'
    )
    await callback.answer()

@dp.callback_query(F.data.startswith("adm_ext30_"))
async def adm_extend_30(callback: CallbackQuery):
    if callback.from_user.id != SUPER_ADMIN_ID:
        await callback.answer("غير مصرح.", show_alert=True); return
    uid = int(callback.data.replace("adm_ext30_", ""))
    new_exp = await database.add_or_update_user(uid, days_to_add=30)
    await database.log_action(uid, 'admin_extend', f'+30 days. New expiry: {new_exp}')
    await callback.message.edit_text(
        f"✅ تم تمديد اشتراك {uid} بـ 30 يوم.\n📅 الانتهاء الجديد: {new_exp[:10]}",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="🔙 رجوع للمستخدمين", callback_data="adm_subs")]
        ])
    )
    await callback.answer()

@dp.callback_query(F.data.startswith("adm_red7_"))
async def adm_reduce_7(callback: CallbackQuery):
    if callback.from_user.id != SUPER_ADMIN_ID:
        await callback.answer("غير مصرح.", show_alert=True); return
    uid = int(callback.data.replace("adm_red7_", ""))
    new_exp = await database.reduce_subscription(uid, 7)
    await database.log_action(uid, 'admin_reduce', f'-7 days. New expiry: {new_exp}')
    await callback.message.edit_text(
        f"✅ تم خصم 7 أيام من اشتراك {uid}.\n📅 الانتهاء الجديد: {new_exp[:10]}",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="🔙 رجوع للمستخدمين", callback_data="adm_subs")]
        ])
    )
    await callback.answer()

@dp.callback_query(F.data.startswith("adm_setexp_"))
async def adm_set_custom_expiry_prompt(callback: CallbackQuery, state: FSMContext):
    if callback.from_user.id != SUPER_ADMIN_ID:
        await callback.answer("غير مصرح.", show_alert=True); return
    uid = int(callback.data.replace("adm_setexp_", ""))
    await state.set_state(MarineStates.admin_waiting_for_custom_expiry)
    await state.update_data(admin_target_uid=uid)
    await callback.message.edit_text(
        f"📅 أدخل تاريخ الانتهاء الجديد للمستخدم <code>{uid}</code>\n"
        "بالصيغة: <b>YYYY-MM-DD</b>\n"
        "مثال: <code>2027-01-15</code>",
        parse_mode='HTML'
    )
    await callback.answer()

@dp.message(MarineStates.admin_waiting_for_custom_expiry, F.text & ~F.text.startswith('/'))
async def adm_set_custom_expiry_input(message: Message, state: FSMContext):
    if message.from_user.id != SUPER_ADMIN_ID:
        return
    data = await state.get_data()
    uid = data.get('admin_target_uid')
    date_str = message.text.strip()
    try:
        iso = f"{date_str}T00:00:00"
        new_exp = await database.set_subscription_expiry(uid, iso)
        await database.log_action(uid, 'admin_set_expiry', f'Exact expiry set to: {new_exp}')
        await state.clear()
        await message.answer(
            f"✅ تم تعيين تاريخ انتهاء {uid} إلى {new_exp[:10]}",
            reply_markup=get_main_menu()
        )
    except ValueError:
        await message.answer("❌ صيغة التاريخ غير صحيحة. أدخل بالصيغة: YYYY-MM-DD")

@dp.callback_query(F.data.startswith("adm_revoke_"))
async def adm_revoke_user(callback: CallbackQuery):
    if callback.from_user.id != SUPER_ADMIN_ID:
        await callback.answer("غير مصرح.", show_alert=True); return
    uid = int(callback.data.replace("adm_revoke_", ""))
    await database.revoke_user(uid)
    await database.log_action(uid, 'admin_revoke', 'Subscription revoked via admin panel')
    await callback.message.edit_text(
        f"🚫 تم إيقاف اشتراك المستخدم {uid}.",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="🔙 رجوع للمستخدمين", callback_data="adm_subs")]
        ])
    )
    await callback.answer()

# ── Audit Log ─────────────────────────────────────────────────
@dp.callback_query(F.data == "adm_audit")
async def adm_audit_list(callback: CallbackQuery):
    if callback.from_user.id != SUPER_ADMIN_ID:
        await callback.answer("غير مصرح.", show_alert=True); return
    users = await database.get_all_users_stats()
    if not users:
        await callback.message.edit_text("لا يوجد مستخدمون.", reply_markup=_admin_main_kb())
        await callback.answer(); return
    buttons = []
    for uid, boat, *_ in users:
        buttons.append([InlineKeyboardButton(text=f"📜 {uid} — {boat}", callback_data=f"adm_log_{uid}")])
    buttons.append([InlineKeyboardButton(text="🔙 رجوع", callback_data="adm_back")])
    await callback.message.edit_text(
        "📜 <b>سجل العمليات</b>\nاختر المستخدم:",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons), parse_mode='HTML'
    )
    await callback.answer()

@dp.callback_query(F.data.startswith("adm_log_"))
async def adm_view_user_log(callback: CallbackQuery):
    if callback.from_user.id != SUPER_ADMIN_ID:
        await callback.answer("غير مصرح.", show_alert=True); return
    import html
    uid = int(callback.data.replace("adm_log_", ""))
    logs = await database.get_user_audit_log(uid, limit=20)
    if not logs:
        await callback.message.edit_text(
            f"لا توجد سجلات للمستخدم {uid}.",
            reply_markup=InlineKeyboardMarkup(inline_keyboard=[
                [InlineKeyboardButton(text="🔙 رجوع", callback_data="adm_audit")]
            ])
        )
        await callback.answer(); return
    msg = f"📜 <b>آخر 20 إجراء للمستخدم</b> <code>{uid}</code>:\n\n"
    for log_id, action_type, details, created_at in logs:
        dt_str = str(created_at)[:19].replace("T", " ")
        det = html.escape(str(details or ""))
        msg += f"🕐 {dt_str}\n🔹 {html.escape(action_type)}: {det}\n\n"
    # Telegram message limit is 4096 chars
    if len(msg) > 4000:
        msg = msg[:3990] + "\n\n... (مقطوع)"
    await callback.message.edit_text(msg, reply_markup=InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="🔙 رجوع", callback_data="adm_audit")]
    ]), parse_mode='HTML')
    await callback.answer()

# ── Back to admin menu ────────────────────────────────────────
@dp.callback_query(F.data == "adm_back")
async def adm_back_to_menu(callback: CallbackQuery):
    if callback.from_user.id != SUPER_ADMIN_ID:
        await callback.answer("غير مصرح.", show_alert=True); return
    await callback.message.edit_text(
        "🔐 <b>لوحة تحكم المشرف — AbsyCode</b>\n\nاختر من القائمة:",
        reply_markup=_admin_main_kb(), parse_mode='HTML'
    )
    await callback.answer()

# ─────────────────────────────────────────────────────────────
# ACTIVATION KEY
# ─────────────────────────────────────────────────────────────
@dp.message(MarineStates.waiting_for_activation_key, F.text & ~F.text.startswith('/'))
async def handle_activation_key(message: Message, state: FSMContext):
    if message.text.strip() == TRIAL_ACTIVATION_KEY:
        now = datetime.now()
        expiry_date = now + timedelta(days=3)
        expiry_iso = expiry_date.isoformat()
        
        # Call database helper to update user status and set subscription expiry date
        await database.add_or_update_user(
            user_id=message.from_user.id,
            days_to_add=3,
            boat_name=f"Guest_{message.from_user.id}",
            expiry_iso=expiry_iso
        )
        await database.log_action(message.from_user.id, 'activation_success', 'TRIAL KEY ACTIVATED (3 days)')
        
        # Clean up tracking set if present
        if message.from_user.id in pending_approval_requests:
            pending_approval_requests.discard(message.from_user.id)
            
        await state.set_state(MarineStates.waiting_for_boat_name)
        await message.answer("مرحباً بك يا ريس! اكتب اسم مركبك أولاً لتفعيل الحساب:")
    else:
        await message.answer("❌ الكود غير صحيح، حاول مرة أخرى.")

@dp.message(MarineStates.waiting_for_boat_name, F.text & ~F.text.startswith('/'))
async def handle_boat_name_input(message: Message, state: FSMContext):
    boat_name = message.text.strip()
    user_id = message.from_user.id
    
    await database.update_boat_name(user_id, boat_name)
    await database.log_action(user_id, 'boat_name_set', f'Boat name initialized to: {boat_name}')
    
    await state.clear()
    success_message = (
        f"✅ تم حفظ اسم مركبك: **{boat_name}**!\n\n"
        "مرحباً بك في AbsyCode Marine Assistant. يمكنك الآن البدء في استخدام البوت وإدارة حسابات مركبك بكل سهولة."
    )
    await message.answer(success_message, reply_markup=get_main_menu(), parse_mode="Markdown")


# ─────────────────────────────────────────────────────────────
# EXISTING MENU HANDLERS
# ─────────────────────────────────────────────────────────────
@dp.message(F.text == "🛳️ مكتبي البحري")
async def handle_marine_office(message: Message, state: FSMContext):
    await state.clear()
    await message.answer("أهلاً بك في مكتبك البحري! 🛳️", reply_markup=get_main_menu())

@dp.message(F.text == "🌊 سجل السرحات")
async def handle_trips_log(message: Message, state: FSMContext):
    await state.clear()
    await message.answer("سجل السرحات مفتوح. 🌊", reply_markup=get_main_menu())

@dp.message(F.text == "🛒 تجهيز رحلة")
async def handle_trip_prep(message: Message, state: FSMContext):
    await state.clear()
    await message.answer("جاهزين للرحلة! 🛒", reply_markup=get_main_menu())

@dp.message(F.text == "🚨 طوارئ وإنقاذ")
async def handle_emergency(message: Message, state: FSMContext):
    await state.clear()
    await message.answer("قسم الطوارئ. 🚨", reply_markup=get_main_menu())

@dp.message(F.text == "💰 الحسابات")
async def handle_finance(message: Message, state: FSMContext):
    await state.clear()
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="➕ إضافة مصروف",  callback_data="finance_add_expense")],
        [InlineKeyboardButton(text="📊 تقارير مالية", callback_data="finance_reports")],
    ])
    await message.answer("قسم المالية 💰", reply_markup=kb)

@dp.callback_query(F.data.startswith("finance_"))
async def process_finance_choice(callback: CallbackQuery, state: FSMContext):
    await callback.message.edit_reply_markup(reply_markup=None)
    if callback.data == "finance_add_expense":
        await state.set_state(MarineStates.waiting_for_expense_details)
        await callback.message.answer("ابعت تفاصيل المصروف أو صورة الفاتورة. ✍️")
    elif callback.data == "finance_reports":
        await cmd_report(callback.message)
    await callback.answer()


@dp.message(F.text == "🌤️ حالة البحر")
async def handle_weather(message: Message):
    """v4.0 — Full daily marine weather forecast with hourly breakdown."""
    user_data = await database.get_user(message.from_user.id)
    boat_name = user_data[2] if user_data and user_data[2] else "مركب غير مسمى"
    wait_msg  = await message.answer("جاري استطلاع حالة البحر على مدار اليوم... 🔭")
    lat, lon  = 27.25, 33.81   # Hurghada / Red Sea default
    api_key   = os.getenv("OPENWEATHER_API_KEY", "")
    try:
        loop = asyncio.get_running_loop()
        # Current weather snapshot
        current_url = (f"https://api.openweathermap.org/data/2.5/weather?lat={lat}&lon={lon}"
                       f"&appid={api_key}&units=metric&lang=ar")
        current_data = await loop.run_in_executor(
            None, lambda: requests.get(current_url, timeout=5).json()
        )
        curr_temp = current_data['main']['temp']
        curr_wind_ms = current_data['wind']['speed']
        curr_wind_knots = curr_wind_ms * 1.94384
        curr_desc = current_data['weather'][0]['description']
        curr_wave_h, curr_wave_desc = _estimate_wave_height(curr_wind_knots)
        curr_safety = _marine_safety(curr_wind_knots)

        # 3-hour forecast for full day breakdown
        forecast_url = (f"https://api.openweathermap.org/data/2.5/forecast?lat={lat}&lon={lon}"
                        f"&appid={api_key}&units=metric&lang=ar")
        forecast_resp = await loop.run_in_executor(
            None, lambda: requests.get(forecast_url, timeout=5).json()
        )

        # Build report header
        raw_text = (f"⚓ تقرير البحر — {boat_name}\n"
                    + "━"*20 + "\n\n"
                    f"📍 الحالة الآن:\n"
                    f"🌡️ {curr_temp:.1f}°C  |  🌬️ {curr_wind_knots:.1f} عقدة\n"
                    f"🌊 ارتفاع الموج: ~{curr_wave_h:.1f}م ({curr_wave_desc})\n"
                    f"☁️ {curr_desc}\n"
                    f"💎 {curr_safety}\n\n"
                    + "━"*20 + "\n"
                    f"📊 التوقعات على مدار اليوم:\n\n")

        # Parse forecast — group today's entries into 4 periods
        today_str = datetime.now().strftime('%Y-%m-%d')
        tomorrow_str = (datetime.now() + timedelta(days=1)).strftime('%Y-%m-%d')
        period_defs = [
            ("الصباح ☀️ (06:00–12:00)", 6, 12),
            ("الظهر 🌤️ (12:00–18:00)", 12, 18),
            ("المساء 🌅 (18:00–00:00)", 18, 24),
            ("الليل 🌙 (00:00–06:00)", 0, 6),
        ]
        periods = {name: [] for name, _, _ in period_defs}

        for entry in forecast_resp.get('list', []):
            dt_txt = entry.get('dt_txt', '')
            is_today = dt_txt.startswith(today_str)
            is_tomorrow_night = (dt_txt.startswith(tomorrow_str) and int(dt_txt[11:13]) < 6)
            if not is_today and not is_tomorrow_night:
                continue
            hour = int(dt_txt[11:13])
            temp = entry['main']['temp']
            wind_knots = entry['wind']['speed'] * 1.94384
            desc = entry['weather'][0]['description']
            wave_h, wave_desc = _estimate_wave_height(wind_knots)
            pd_entry = {'time': dt_txt[11:16], 'temp': temp,
                        'wind_knots': wind_knots, 'desc': desc,
                        'wave_h': wave_h, 'wave_desc': wave_desc}
            for p_name, p_start, p_end in period_defs:
                if p_start <= hour < p_end:
                    periods[p_name].append(pd_entry)
                    break

        for p_name, _, _ in period_defs:
            p_entries = periods[p_name]
            if not p_entries:
                continue
            raw_text += f"🕐 {p_name}:\n"
            for e in p_entries:
                raw_text += (f"   ⏰ {e['time']} → "
                            f"🌡️{e['temp']:.0f}°C | "
                            f"🌬️{e['wind_knots']:.0f}kn | "
                            f"🌊{e['wave_h']:.1f}م | "
                            f"{e['desc']}\n")
            max_wind = max(e['wind_knots'] for e in p_entries)
            raw_text += f"   {_marine_safety(max_wind)}\n\n"

        # AI formatting via Gemini
        final_text = raw_text
        try:
            gemini_key = os.getenv("GEMINI_API_KEY")
            if gemini_key:
                ai_url = f"https://generativelanguage.googleapis.com/v1/models/gemini-2.5-flash:generateContent?key={gemini_key}"
                prompt = (f"قم بصياغة تقرير الطقس البحري هذا بأسلوب بحري مصري ودود ومختصر للقبطان. "
                         f"حافظ على كل الأرقام والفترات الزمنية وبيانات الرياح والموج كما هي:\n{raw_text}")
                payload = {"contents": [{"parts": [{"text": prompt}]}]}
                resp = await loop.run_in_executor(
                    None, lambda: requests.post(ai_url, json=payload, timeout=10)
                )
                if resp.status_code == 200:
                    ai_data = resp.json()
                    if "candidates" in ai_data and len(ai_data["candidates"]) > 0:
                        final_text = ai_data["candidates"][0]["content"]["parts"][0]["text"].strip()
        except Exception as ai_err:
            logging.warning(f"Gemini weather summary failed: {ai_err}")

        if len(final_text) > 4000:
            final_text = final_text[:3990] + "\n\n... (مقطوع)"
        await wait_msg.edit_text(final_text)
    except Exception as e:
        logging.error(f"Weather error: {e}")
        traceback.print_exc()
        await wait_msg.edit_text("مش قادر أوصل لبيانات الطقس حالياً. 😅")


def _estimate_wave_height(wind_knots: float) -> tuple:
    """Beaufort scale wind-to-wave height approximation for marine reports."""
    if wind_knots < 1:    return 0.0, "هادئ تماماً 🏖️"
    elif wind_knots < 4:  return 0.1, "هادئ 🏖️"
    elif wind_knots < 7:  return 0.3, "أمواج خفيفة 🌊"
    elif wind_knots < 11: return 0.6, "أمواج معتدلة 🌊"
    elif wind_knots < 17: return 1.2, "أمواج متوسطة ⚠️"
    elif wind_knots < 22: return 2.5, "أمواج عالية ⚠️"
    elif wind_knots < 28: return 4.0, "أمواج عالية جداً 🚨"
    else:                 return 6.0, "بحر هائج 🚨🚨"


def _marine_safety(wind_knots: float) -> str:
    """Marine trip safety assessment based on wind speed."""
    if wind_knots < 10:   return "✅ آمن للإبحار — البحر هادي ومثالي للرحلات"
    elif wind_knots < 17: return "⚠️ مناسب بحذر — الرياح معتدلة"
    elif wind_knots < 22: return "⚠️ غير مستحب — الرياح قوية"
    else:                 return "🚫 خطر — يُنصح بعدم الإبحار"


# ─────────────────────────────────────────────────────────────
# ADMIN APPROVAL CALLBACKS
# ─────────────────────────────────────────────────────────────
@dp.callback_query(F.data.startswith("approve_"))
async def admin_approve_request(callback: CallbackQuery):
    if callback.from_user.id != SUPER_ADMIN_ID:
        await callback.answer("غير مصرح.", show_alert=True); return
    user_id = callback.data.split("_")[1]
    kb = InlineKeyboardMarkup(inline_keyboard=[[
        InlineKeyboardButton(text="أسبوع",  callback_data=f"dur_7_{user_id}"),
        InlineKeyboardButton(text="شهر",    callback_data=f"dur_30_{user_id}"),
        InlineKeyboardButton(text="سنة",    callback_data=f"dur_365_{user_id}"),
    ]])
    await callback.message.edit_text(callback.message.text + "\n\nاختر مدة الاشتراك:", reply_markup=kb)
    await callback.answer()

@dp.callback_query(F.data.startswith("reject_"))
async def admin_reject_request(callback: CallbackQuery):
    if callback.from_user.id != SUPER_ADMIN_ID:
        await callback.answer("غير مصرح.", show_alert=True); return
    user_id = int(callback.data.split("_")[1])
    try: await bot.send_message(user_id, "عذراً، لم يتم قبول طلب انضمامك حالياً.")
    except Exception: pass
    await callback.message.edit_text(callback.message.text + "\n\n❌ تم الرفض.", reply_markup=None)
    pending_approval_requests.discard(user_id)
    await callback.answer()

@dp.callback_query(F.data.startswith("dur_"))
async def admin_set_duration(callback: CallbackQuery):
    if callback.from_user.id != SUPER_ADMIN_ID:
        await callback.answer("غير مصرح.", show_alert=True); return
    parts   = callback.data.split("_")
    days    = int(parts[1])
    user_id = int(parts[2])
    await database.add_or_update_user(user_id, days)
    dur_txt = "أسبوع" if days == 7 else "شهر" if days == 30 else "سنة"
    try: await bot.send_message(user_id, f"تم تفعيل اشتراكك لمدة {dur_txt}! 🛥️✨\nاكتب /start للبدء.")
    except Exception: pass
    await callback.message.edit_text(
        callback.message.text.split("\n\nاختر")[0] + f"\n\n✅ تم التفعيل لمدة {dur_txt}.", reply_markup=None)
    pending_approval_requests.discard(user_id)
    await callback.answer()

# ─────────────────────────────────────────────────────────────
# ADMIN COMMANDS
# ─────────────────────────────────────────────────────────────
@dp.message(Command("add_user"))
async def cmd_add_user(message: Message):
    if message.from_user.id != SUPER_ADMIN_ID: return
    args = message.text.split()
    if len(args) != 3:
        await message.answer("الاستخدام: /add_user <user_id> <days>"); return
    try:
        expiry = await database.add_or_update_user(int(args[1]), int(args[2]))
        await message.answer(f"✅ تم تفعيل {args[1]} حتى {expiry[:10]}")
    except ValueError:
        await message.answer("أرقام غير صحيحة.")

@dp.message(Command("stats"))
async def cmd_stats(message: Message):
    import html
    if message.from_user.id != SUPER_ADMIN_ID:
        await cmd_report(message); return
    users = await database.get_all_users_stats()
    if not users:
        await message.answer("لا يوجد مستخدمين."); return
    msg = "📊 <b>إحصائيات المستخدمين:</b>\n\n"
    for uid, b_name, is_act, expiry, reg_at in users:
        status   = "✅ نشط" if is_act else "❌ موقوف"
        exp_date = str(expiry)[:10]  if expiry  else "N/A"
        reg_date = str(reg_at)[:10]  if reg_at  else "N/A"
        msg += f"👤 <code>{uid}</code> | 🛥️ {html.escape(b_name)}\n⏳ {status} حتى {exp_date} | تسجيل: {reg_date}\n\n"
    await message.answer(msg, parse_mode='HTML')

@dp.message(Command("set_name"))
async def cmd_set_name(message: Message):
    args = message.text.split(maxsplit=1)
    if len(args) < 2:
        await message.answer("الاستخدام: /set_name <اسم المركب>"); return
    await database.update_boat_name(message.from_user.id, args[1].strip())
    await message.answer(f"✅ اسم مركبك: {args[1].strip()} 🛥️")

@dp.message(Command("last"))
async def cmd_last(message: Message):
    import html
    invoice, items = await database.get_last_invoice(message.from_user.id)
    if not invoice:
        await message.answer("لا توجد فواتير مسجلة."); return
    inv_id, created_at, amount, currency, category, raw_text, image_file_id = invoice
    msg = (f"📄 <b>آخر فاتورة #{inv_id}</b>\n"
           f"📅 {str(created_at)[:10]}\n"
           f"🏷️ {html.escape(category or 'غير مصنف')}\n"
           f"💰 {amount:,.2f} {html.escape(currency)}\n"
           f"📸 صورة: {'✅' if image_file_id else '❌'}\n")
    if items:
        msg += "\n📦 الأصناف:\n" + "\n".join(f"🔹 {html.escape(str(n))} ({q}x) — {p:,.2f}" for n, q, p in items)
    await message.answer(msg, parse_mode='HTML')

@dp.message(Command("delete_last"))
async def cmd_delete_last(message: Message):
    success = await database.delete_last_record(message.from_user.id)
    await message.answer("✅ تم مسح آخر فاتورة." if success else "❌ لا توجد فواتير.")

@dp.message(Command("help"))
async def cmd_help(message: Message):
    await message.answer(
        "❓ <b>كيفية الاستخدام:</b>\n\n"
        "📸 صورة فاتورة للاستخراج التلقائي\n"
        "✍️ نص مثل <code>500 بنزين</code> للتسجيل الفوري\n"
        "🎧 رسالة صوتية لتسجيل مصروف\n\n"
        "🔹 <b>الأوامر:</b>\n"
        "/start /report /export /last /delete_last /set_name /help",
        parse_mode='HTML', reply_markup=get_main_menu()
    )

# ─────────────────────────────────────────────────────────────
# GENERIC TEXT HANDLER
# ─────────────────────────────────────────────────────────────
@dp.message(F.text & ~F.text.startswith('/'))
async def handle_text(message: Message, state: FSMContext):
    text          = message.text.strip()
    text_lower    = text.lower()
    user_id       = message.from_user.id
    user_data     = await database.get_user(user_id)
    boat_name     = user_data[2] if user_data and user_data[2] else "AbsyCode"
    current_state = await state.get_state()
    # Guard: let other FSM handlers deal with these states
    if current_state in [
        MarineStates.waiting_for_activation_key.state,
        MarineStates.waiting_for_boat_name.state,
        MarineStates.waiting_for_invoice_details.state,
        MarineStates.waiting_for_edit_value.state,
        MarineStates.waiting_for_image_invoice_sel.state,
        MarineStates.waiting_for_image_upload.state,
        MarineStates.waiting_for_invoice_category.state,
        MarineStates.admin_waiting_for_custom_expiry.state,
    ]:
        return
    # Delete shortcut
    if any(k in text_lower for k in ['امسح', 'احذف', 'delete']):
        success = await database.delete_last_record(user_id)
        await message.answer("✅ تم مسح آخر فاتورة!" if success else "مفيش حاجة امسحها!")
        return
    # Instant keyword detection
    category = None
    if 'بنزين' in text_lower or 'سولار' in text_lower:   category = 'بنزين'
    elif 'ماركت' in text_lower or 'أكل' in text_lower:   category = 'ماركت'
    elif 'صيانة' in text_lower or 'تصليح' in text_lower: category = 'صيانة'
    elif 'إكرامية' in text_lower or 'تيبس' in text_lower:category = 'إكرامية'
    elif 'نظافة' in text_lower:                           category = 'أدوات نظافة'
    if category:
        amount = _extract_amount(text)
        await database.save_invoice(user_id, amount, "جنيه", text, category=category)
        await message.answer(f"✅ {amount:,.2f} جنيه — {category} لمركب {boat_name}! 🫡", reply_markup=get_main_menu())
        await state.clear()
        return
    # AI fallback
    ai_msg = await message.answer("🤖 جاري فهم طلبك...")
    try:
        if current_state == MarineStates.waiting_for_expense_details.state:
            amount    = _extract_amount(text)
            category  = "عام"
            ai_intent = analyze_with_ai(text, mode='text_intent')
            if ai_intent:
                amount   = ai_intent.get('amount', amount)
                category = ai_intent.get('category', category)
            await database.save_invoice(user_id, amount, "جنيه", text, category=category)
            await ai_msg.edit_text(f"✅ {amount:,.2f} جنيه في {category} لمركب {boat_name}! 🫡")
        else:
            ai_intent = analyze_with_ai(text, mode='text_intent', boat_name=boat_name)
            action    = ai_intent.get('action') if ai_intent else None
            amount    = ai_intent.get('amount', 0) if ai_intent else 0
            category  = ai_intent.get('category', 'عام') if ai_intent else 'عام'
            if action in ['expense', 'add'] and amount > 0:
                await database.save_invoice(user_id, amount, "جنيه", text, category=category)
                await ai_msg.edit_text(f"✅ {amount:,.2f} جنيه في {category} لمركب {boat_name}! 🫡")
            elif action == 'delete_category':
                ok = await database.delete_category(user_id, category)
                await ai_msg.edit_text(f"✅ تم تنظيف قسم {category}!" if ok else f"قسم {category} فاضي أصلاً!")
            elif action == 'query':
                await ai_msg.delete()
                await cmd_report(message)
            else:
                resp = analyze_with_ai(text, mode='chat', boat_name=boat_name)
                await ai_msg.edit_text(resp or "مش فاهم حضرتك، حاول تاني.")
    except Exception as e:
        logging.error(f"Text handler error: {e}")
        amount = _extract_amount(text)
        if amount > 0:
            await database.save_invoice(user_id, amount, "جنيه", text, category="عام")
            await ai_msg.edit_text(f"✅ {amount:,.2f} جنيه في قسم عام!")
        else:
            await ai_msg.edit_text("عذراً، لم أستطع فهم الطلب.")
    finally:
        await state.clear()


# ─────────────────────────────────────────────────────────────
# STARTUP
# ─────────────────────────────────────────────────────────────
async def setup_bot_commands(bot: Bot):
    await bot.set_my_commands([
        BotCommand(command="start",       description="بدء البوت"),
        BotCommand(command="report",      description="تقرير المصاريف"),
        BotCommand(command="export",      description="تحميل Excel و Word"),
        BotCommand(command="last",        description="آخر فاتورة"),
        BotCommand(command="delete_last", description="مسح آخر فاتورة"),
        BotCommand(command="set_name",    description="تغيير اسم المركب"),
        BotCommand(command="admin",       description="لوحة التحكم (مشرف)"),
        BotCommand(command="help",        description="المساعدة"),
    ])

async def main():
    await database.init_db()
    logging.info("Starting AbsyCode Marine Bot v3.0...")
    await setup_bot_commands(bot)
    await dp.start_polling(bot)

if __name__ == "__main__":
    asyncio.run(main())
